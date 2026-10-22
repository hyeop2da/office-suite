# 시험용 XLSX를 openpyxl로 만든다(직접 생성한 시료. Excel이 만든 파일이 아님).
# 사용: python tools/make_xlsx_fixtures.py fixtures/xlsx
import os, sys, zipfile, datetime
from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

out = sys.argv[1]
os.makedirs(out, exist_ok=True)

# 1) types.xlsx: 숫자·문자열·불리언·날짜·수식·공유 문자열·여러 시트·숨김 시트
wb = Workbook()
ws = wb.active
ws.title = 'Sheet1'
ws['A1'] = '정수'; ws['B1'] = 42
ws['A2'] = '실수'; ws['B2'] = 3.14159
ws['A3'] = '음수'; ws['B3'] = -7
ws['A4'] = '지수'; ws['B4'] = 1.5e-7
ws['A5'] = '큰 수'; ws['B5'] = 12345678901
ws['A6'] = '퍼센트'; ws['B6'] = 0.25; ws['B6'].number_format = '0%'
ws['A7'] = '날짜'; ws['B7'] = datetime.date(2026, 10, 22); ws['B7'].number_format = 'yyyy-mm-dd'
ws['A8'] = '불리언'; ws['B8'] = True
ws['A9'] = '합계'; ws['B9'] = '=SUM(B1:B3)'
ws['A10'] = '특수 & <문자> "따옴표"'; ws['B10'] = '  앞뒤 공백  '
ws['A11'] = '같은말'; ws['B11'] = '같은말'            # 같은 공유 문자열을 두 셀이 쓴다
ws['A12'] = '줄\n바꿈'
ws['D5'] = '떨어진 칸'                                   # 비어 있는 칸 건너뜀
ws2 = wb.create_sheet('두 번째 & 시트')
ws2['A1'] = '둘째 시트'; ws2['B1'] = 100
ws3 = wb.create_sheet('숨김')
ws3['A1'] = '숨겨진 시트'
ws3.sheet_state = 'hidden'
wb.save(f'{out}/types.xlsx')

# 2) rich.xlsx: 서식이 섞인 공유 문자열
wb = Workbook()
ws = wb.active
ws['A1'] = CellRichText('굵게', TextBlock(InlineFont(i=True), ' 기울임'), ' 보통')
ws['A2'] = '일반 문자열'
wb.save(f'{out}/rich.xlsx')

# 이 환경의 openpyxl은 문자열을 inlineStr로 쓴다. Excel처럼 공유 문자열(t="s")을 쓰는 변형도 만든다.
import re
def to_shared(src, dst):
    zin = zipfile.ZipFile(src)
    table, index = [], {}
    def conv(xml):
        def rep(m):
            attrs, content = m.group(1), m.group(2)
            if content not in index:
                index[content] = len(table); table.append(content)
            attrs = attrs.replace(' t="inlineStr"', '')
            return f'<c{attrs} t="s"><v>{index[content]}</v></c>'
        return re.sub(r'<c((?:\s[^>]*?)? t="inlineStr"[^>]*?)><is>(.*?)</is></c>', rep, xml, flags=re.S)
    parts = {}
    total = 0
    for item in zin.infolist():
        data = zin.read(item.filename)
        if re.match(r'xl/worksheets/sheet\d+\.xml$', item.filename):
            before = data.decode('utf8')
            after = conv(before)
            total += after.count(' t="s"')
            data = after.encode('utf8')
        parts[item.filename] = (item, data)
    sst = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
           f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{total}" uniqueCount="{len(table)}">'
           + ''.join(f'<si>{c}</si>' for c in table) + '</sst>')
    ct = parts['[Content_Types].xml'][1].decode('utf8')
    ct = ct.replace('</Types>', '<Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
    rels = parts['xl/_rels/workbook.xml.rels'][1].decode('utf8')
    rels = rels.replace('</Relationships>', '<Relationship Id="rIdSst" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/></Relationships>')
    zout = zipfile.ZipFile(dst, 'w', zipfile.ZIP_DEFLATED)
    for name, (item, data) in parts.items():
        if name == '[Content_Types].xml': data = ct.encode('utf8')
        if name == 'xl/_rels/workbook.xml.rels': data = rels.encode('utf8')
        zout.writestr(item, data)
    zout.writestr('xl/sharedStrings.xml', sst)
    zout.close(); zin.close()

import shutil
shutil.copy(f'{out}/types.xlsx', f'{out}/types-inline.xlsx')
to_shared(f'{out}/types-inline.xlsx', f'{out}/types.xlsx')
to_shared(f'{out}/rich.xlsx', f'{out}/_rich_shared.xlsx')
os.replace(f'{out}/_rich_shared.xlsx', f'{out}/rich.xlsx')
print('ok', sorted(os.listdir(out)))
