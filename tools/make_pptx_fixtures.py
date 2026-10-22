# 시험용 PPTX를 python-pptx로 만든다(직접 생성한 시료. PowerPoint가 만든 파일이 아님).
# 사용: python tools/make_pptx_fixtures.py fixtures/pptx
import os, sys, re, zipfile
from pptx import Presentation
from pptx.util import Inches, Pt

out = sys.argv[1]
os.makedirs(out, exist_ok=True)

prs = Presentation()
# 1: 제목 + 본문(여러 글자 모양) + 슬라이드 번호 필드
s = prs.slides.add_slide(prs.slide_layouts[1])
s.shapes.title.text = '제목 슬라이드 & <시험>'
body = s.placeholders[1].text_frame
body.text = '첫째 줄'
p = body.add_paragraph()
r = p.add_run(); r.text = '굵게'; r.font.bold = True
r2 = p.add_run(); r2.text = ' 보통 글자'
p = body.add_paragraph(); p.text = '  앞뒤 공백이 있는 줄  '
# 2: 표
s2 = prs.slides.add_slide(prs.slide_layouts[5])
s2.shapes.title.text = '표 슬라이드'
tbl = s2.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1.5)).table
for i in range(2):
    for j in range(2):
        tbl.cell(i, j).text = f'셀 {i+1}-{j+1}'
# 3: 텍스트 상자 + 발표자 노트
s3 = prs.slides.add_slide(prs.slide_layouts[6])
tb = s3.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
tb.text_frame.text = '텍스트 상자 😀 漢字'
s3.notes_slide.notes_text_frame.text = '발표자 노트 문장'
prs.save(f'{out}/basic.pptx')

# 필드(슬라이드 번호)를 직접 넣은 변형: 첫 슬라이드 본문 끝에 a:fld 문단
tmp = f'{out}/_src.pptx'
os.replace(f'{out}/basic.pptx', tmp)
zin = zipfile.ZipFile(tmp)
zout = zipfile.ZipFile(f'{out}/basic.pptx', 'w', zipfile.ZIP_DEFLATED)
fld = '<a:p><a:r><a:rPr lang="ko-KR"/><a:t>번호: </a:t></a:r><a:fld id="{B6F15528-21DE-4FAA-801E-634DDDAF4B2B}" type="slidenum"><a:rPr lang="ko-KR"/><a:t>1</a:t></a:fld></a:p>'
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == 'ppt/slides/slide1.xml':
        x = data.decode('utf8')
        i = x.rfind('</p:txBody>')
        x = x[:i] + fld + x[i:]
        data = x.encode('utf8')
    zout.writestr(item, data)
zout.close(); zin.close(); os.remove(tmp)
print('ok', os.listdir(out))
