# 시험용 DOCX를 python-docx로 만든다(직접 만든 시료. 실제 Word로 만든 파일이 아님).
# 사용: python tools/make_docx_fixtures.py fixtures/docx
import sys, zipfile, shutil, os
from docx import Document

out = sys.argv[1]
os.makedirs(out, exist_ok=True)

d = Document()
d.add_paragraph('안녕하세요 & <테스트> "따옴표"')
d.add_paragraph('  앞뒤 공백이 있는 문단  ')
p = d.add_paragraph()
r = p.add_run('굵게'); r.bold = True
p.add_run(' 보통 글자')
d.add_paragraph('')
d.add_paragraph('이모지 😀 와 한자 漢字')
d.save(f'{out}/basic.docx')

d = Document()
d.add_paragraph('표 앞 문단')
t = d.add_table(rows=2, cols=2)
for i in range(2):
    for j in range(2):
        t.cell(i, j).text = f'셀 {i+1}-{j+1}'
d.add_paragraph('표 뒤 문단')
d.save(f'{out}/table.docx')

d = Document()
d.add_paragraph('본문')
sec = d.sections[0]
sec.header.paragraphs[0].text = '머리글 텍스트'
sec.footer.paragraphs[0].text = '바닥글 텍스트'
d.save(f'{out}/headerfooter.docx')

# 텍스트 상자(mc:AlternateContent: Choice와 Fallback에 같은 글자)를 직접 넣는다
tmp = f'{out}/_tb.docx'
d = Document()
d.add_paragraph('상자 앞')
d.save(tmp)
zin = zipfile.ZipFile(tmp)
zout = zipfile.ZipFile(f'{out}/textbox.docx', 'w', zipfile.ZIP_DEFLATED)
NS = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
      'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
      'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
      'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape" '
      'xmlns:v="urn:schemas-microsoft-com:vml" xmlns:o="urn:schemas-microsoft-com:office:office"')
# Word 2010 이후가 쓰는 구조: Choice(wps 앵커 도형 + txbxContent) / Fallback(VML 글상자 + txbxContent)
tb = ('<w:p><w:r><mc:AlternateContent ' + NS + '>'
      '<mc:Choice Requires="wps"><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
      '<wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:posOffset>0</wp:posOffset></wp:positionH>'
      '<wp:positionV relativeFrom="paragraph"><wp:posOffset>0</wp:posOffset></wp:positionV>'
      '<wp:extent cx="2000000" cy="600000"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapSquare wrapText="bothSides"/>'
      '<wp:docPr id="1" name="TextBox 1"/><wp:cNvGraphicFramePr/>'
      '<a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><wps:wsp><wps:cNvSpPr txBox="1"/>'
      '<wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="2000000" cy="600000"/></a:xfrm><a:prstGeom prst="rect"><a:avLst/></a:prstGeom></wps:spPr>'
      '<wps:txbx><w:txbxContent><w:p><w:r><w:t>상자 안 글자</w:t></w:r></w:p></w:txbxContent></wps:txbx><wps:bodyPr/></wps:wsp></a:graphicData></a:graphic>'
      '</wp:anchor></w:drawing></mc:Choice>'
      '<mc:Fallback><w:pict><v:shapetype id="_x0000_t202" coordsize="21600,21600" o:spt="202" path="m,l,21600r21600,l21600,xe"><v:stroke joinstyle="miter"/><v:path gradientshapeok="t" o:connecttype="rect"/></v:shapetype>'
      '<v:shape id="TextBox1" o:spid="_x0000_s1026" type="#_x0000_t202" style="position:absolute;width:157.5pt;height:47.25pt"><v:textbox><w:txbxContent><w:p><w:r><w:t>상자 안 글자</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback>'
      '</mc:AlternateContent></w:r></w:p>')
for item in zin.infolist():
    data = zin.read(item.filename)
    if item.filename == 'word/document.xml':
        s = data.decode('utf8')
        s = s.replace('</w:body>', tb + '</w:body>', 1) if s.count('<w:sectPr') == 0 else s.replace('<w:sectPr', tb + '<w:sectPr', 1)
        data = s.encode('utf8')
    zout.writestr(item, data)
zout.close(); zin.close(); os.remove(tmp)
print('ok', os.listdir(out))
